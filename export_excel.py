"""Exporta a Excel los datos actuales de la base de datos (mismos datos que el dashboard).

Script de uso manual desde el equipo local, nada del stack lo llama. Necesita openpyxl,
que no esta en requirements.txt a proposito para no meter la dependencia en la imagen.
"""

import argparse
import os
from datetime import datetime

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.db import Database

load_dotenv()


def get_db() -> Database:
    return Database(
        host=os.getenv("DB_HOST", "localhost"),
        user=os.getenv("DB_USER", "root"),
        password=os.getenv("DB_PASSWORD", ""),
        database=os.getenv("DB_NAME", "competitor_monitor"),
        port=int(os.getenv("DB_PORT", "3306")),
    )


def filter_by_competitor(rows: list[dict], competitor: str) -> list[dict]:
    """Filtra filas por nombre de competidor. Las filas usan la clave
    'competitor' salvo el resumen de competidores, que usa 'name'."""
    return [row for row in rows if row.get("competitor", row.get("name")) == competitor]


def write_sheet(wb: Workbook, title: str, rows: list[dict]):
    ws = wb.create_sheet(title=title)
    if not rows:
        ws.append(["Sin datos"])
        return

    headers = list(rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([row.get(h) for h in headers])

    for i, header in enumerate(headers, start=1):
        max_len = max(
            [len(str(header))] + [len(str(row.get(header, ""))) for row in rows]
        )
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 50)

    ws.freeze_panes = "A2"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--competitor",
        help="Nombre exacto del competidor (columna 'name' en la tabla competitors). "
             "Si se omite, exporta todos los competidores como hasta ahora.",
    )
    args = parser.parse_args()

    db = get_db()

    sheets = {
        "Catalogo actual": db.get_latest_snapshots(),
        "Resumen competidores": db.get_competitor_stats(),
        "Eventos de precio": db.get_recent_price_events(hours=24),
        "Cambios disponibilidad": db.get_recent_availability_events(hours=24),
        "Productos nuevos": db.get_recently_added_products(hours=24),
        "Productos eliminados": db.get_recently_removed_products(hours=24),
    }

    slug = "competidores"
    if args.competitor:
        sheets = {title: filter_by_competitor(rows, args.competitor) for title, rows in sheets.items()}
        slug = args.competitor.lower().replace(" ", "_")

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        write_sheet(wb, title, rows)

    out_dir = "exports"
    os.makedirs(out_dir, exist_ok=True)
    filename = os.path.join(out_dir, f"{slug}_{datetime.now():%Y%m%d_%H%M}.xlsx")
    wb.save(filename)
    print(f"Excel generado: {filename}")


if __name__ == "__main__":
    main()
