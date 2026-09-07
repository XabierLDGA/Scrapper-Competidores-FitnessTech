"""Convierte el Excel de emparejamiento de producto en el SQL que puebla
`titanium_pairs`.

Script de uso manual desde el equipo local, nada del stack lo llama. Necesita
openpyxl, que no esta en requirements.txt a proposito para no meter la
dependencia en la imagen (igual que export_excel.py).

Emite SQL en vez de escribir en la base por dos motivos que no se pueden
esquivar: openpyxl no existe dentro del contenedor, y el contenedor `mysql`
no publica el 3306 fuera de la red interna de Docker. De regalo, el
emparejamiento queda versionado como texto legible y el diff entre dos
revisiones de producto se lee.

    python import_comparativa.py
    scp data/titanium_pairs.sql deploy@168.119.241.200:/tmp/
    ssh deploy@168.119.241.200 \\
      'docker exec -i mysql sh -c \\
         '"'"'mysql -u root -p"$MYSQL_ROOT_PASSWORD" competitor_monitor'"'"' \\
       < /tmp/titanium_pairs.sql'

Ojo a la contrasena: se expande DENTRO del contenedor, que es el unico sitio
donde existe. Y va con root y no con el usuario de la aplicacion (`scraper`,
el de DB_USER), que no tiene permisos para crear tablas.
"""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path("data/comparativa-titanium.xlsx")
SALIDA = Path("data/titanium_pairs.sql")

# Rotulo de la cabecera -> columna de la tabla. Lo que no este aqui se
# ignora: las columnas de precio y diferencia del Excel son una foto del dia
# en que producto lo monto, y los precios buenos salen del scrapper.
COLUMNAS = {
    "SKU Fitness Tech": "ft_sku",
    "Máquina Fitness Tech": "ft_title",
    "Equivalencia": "equivalencia",
    "Máquina Titanium": "titanium_title",
    "URL Titanium": "titanium_url",
    "Observaciones": "observaciones",
}

CAMPOS = ["gama", "orden", "ft_sku", "ft_title", "equivalencia",
          "titanium_title", "titanium_url", "observaciones"]


def _texto(valor):
    """Celda -> str limpio, o None si esta vacia. Las celdas del Excel llegan
    con espacios de sobra y saltos de linea de cuando se pego el texto."""
    if valor is None:
        return None
    limpio = str(valor).strip()
    return limpio or None


def leer_pares(ruta) -> list[dict]:
    """Los pares del Excel, una hoja por gama, en el orden en que los dejo
    producto.

    Lanza ValueError si falta una columna o si un SKU sale dos veces: son
    errores del fichero de producto y hay que verlos, no resolverlos por el.
    """
    libro = load_workbook(ruta, data_only=True)
    pares = []
    vistos = {}

    for hoja in libro.worksheets:
        filas = hoja.iter_rows(values_only=True)
        cabecera = next(filas, None)
        if cabecera is None:
            continue

        # Por rotulo y no por posicion: anadir una columna al Excel no debe
        # romper el importador.
        indice = {}
        for i, celda in enumerate(cabecera):
            rotulo = _texto(celda)
            if rotulo in COLUMNAS:
                indice[COLUMNAS[rotulo]] = i

        faltan = [r for r, campo in COLUMNAS.items() if campo not in indice]
        if faltan:
            raise ValueError(
                f"En la hoja '{hoja.title}' faltan columnas: {', '.join(faltan)}"
            )

        orden = 0
        for fila in filas:
            sku = _texto(fila[indice["ft_sku"]]) if indice["ft_sku"] < len(fila) else None
            if not sku:
                continue  # fila en blanco del final de la hoja

            if sku in vistos:
                raise ValueError(
                    f"El SKU {sku} sale dos veces: en '{vistos[sku]}' y en "
                    f"'{hoja.title}'. Corrigelo en el Excel de producto."
                )
            vistos[sku] = hoja.title

            par = {"gama": hoja.title, "orden": orden, "ft_sku": sku}
            for campo, i in indice.items():
                if campo != "ft_sku":
                    par[campo] = _texto(fila[i]) if i < len(fila) else None
            pares.append(par)
            orden += 1

    return pares


def _valor(v) -> str:
    """Un campo listo para meter en el INSERT."""
    if v is None:
        return "NULL"
    if isinstance(v, int):
        return str(v)
    return "'" + str(v).replace("\\", "\\\\").replace("'", "\\'") + "'"


# El informe va dentro del propio .sql, detras del COMMIT: asi se calcula
# donde estan los datos y sale por pantalla al aplicarlo, sin herramienta
# aparte. Nada de esto corrige nada; son avisos para producto.
INFORME = """
-- ---------------------------------------------------------------------
-- Informe. Nada de lo que salga aqui es un fallo del sistema: son cosas
-- que producto tiene que mirar en su fichero.
-- ---------------------------------------------------------------------

SELECT '1) URLs de Titanium que no existen en el catalogo vigilado' AS informe;
SELECT tp.gama, tp.ft_sku, tp.titanium_url
FROM titanium_pairs tp
LEFT JOIN products p
       ON p.url = tp.titanium_url
      AND p.competitor_id = (SELECT id FROM competitors WHERE name = 'Titanium Strength')
WHERE tp.titanium_url IS NOT NULL AND p.id IS NULL;

SELECT '2) SKUs nuestros que no estan publicados en Fitness Tech ES' AS informe;
SELECT tp.gama, tp.ft_sku, tp.ft_title
FROM titanium_pairs tp
LEFT JOIN products p
       ON p.sku = tp.ft_sku
      AND p.status = 'active'
      AND p.competitor_id = (SELECT id FROM competitors WHERE name = 'Fitness Tech')
WHERE p.id IS NULL;

SELECT '3) De esos, cuales existen con otro prefijo (SE- / PSE-)' AS informe;
SELECT tp.ft_sku AS en_el_excel, p.sku AS en_la_tienda, p.title
FROM titanium_pairs tp
JOIN products p
  ON p.sku <> tp.ft_sku
 AND SUBSTRING_INDEX(p.sku, '-', -1) = SUBSTRING_INDEX(tp.ft_sku, '-', -1)
 AND p.status = 'active'
 AND p.competitor_id = (SELECT id FROM competitors WHERE name = 'Fitness Tech')
LEFT JOIN products exacto
       ON exacto.sku = tp.ft_sku
      AND exacto.status = 'active'
      AND exacto.competitor_id = (SELECT id FROM competitors WHERE name = 'Fitness Tech')
WHERE exacto.id IS NULL;
"""


def generar_sql(pares: list[dict]) -> str:
    """El reemplazo completo, envuelto en una transaccion, mas el informe.

    Reemplazo y no upsert: el Excel es la verdad entera, asi que si producto
    quita una fila tiene que desaparecer tambien aqui. Y la transaccion
    garantiza que la tabla no se quede a medias si el fichero viene roto.
    """
    lineas = [
        "-- Generado por import_comparativa.py a partir de",
        "-- data/comparativa-titanium.xlsx. No editar a mano.",
        "",
        "USE competitor_monitor;",
        "",
        "START TRANSACTION;",
        "",
        "DELETE FROM titanium_pairs;",
        "",
    ]

    for par in pares:
        valores = ", ".join(_valor(par.get(c)) for c in CAMPOS)
        lineas.append(
            f"INSERT INTO titanium_pairs ({', '.join(CAMPOS)}) VALUES ({valores});"
        )

    lineas += ["", "COMMIT;", "", INFORME]
    return "\n".join(lineas) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", nargs="?", default=XLSX, type=Path,
                        help=f"Excel de producto (por defecto {XLSX})")
    parser.add_argument("-o", "--salida", default=SALIDA, type=Path,
                        help=f"Fichero SQL a escribir (por defecto {SALIDA})")
    args = parser.parse_args()

    try:
        pares = leer_pares(args.xlsx)
    except ValueError as e:
        # Sin escribir el .sql: mas vale no tener fichero que tener uno malo.
        print(f"Error en el Excel: {e}", file=sys.stderr)
        return 1

    args.salida.parent.mkdir(parents=True, exist_ok=True)
    args.salida.write_text(generar_sql(pares), encoding="utf-8")

    gamas = {p["gama"] for p in pares}
    print(f"{len(pares)} pares en {len(gamas)} gamas -> {args.salida}")
    for gama in sorted(gamas):
        print(f"  {gama}: {sum(1 for p in pares if p['gama'] == gama)}")
    # La contrasena se expande dentro del contenedor: fuera no existe.
    remoto = f"/tmp/{args.salida.name}"
    print("\nAplicalo con:")
    print(f"  scp {args.salida} deploy@168.119.241.200:/tmp/")
    print("  ssh deploy@168.119.241.200 'docker exec -i mysql sh -c "
          "'\"'\"'mysql -u root -p\"$MYSQL_ROOT_PASSWORD\" competitor_monitor'\"'\"' "
          f"< {remoto}'")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
